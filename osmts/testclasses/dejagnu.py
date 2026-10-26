import re
import shutil
import subprocess
from pathlib import Path
from openpyxl import Workbook

from .errors import GitCloneError, RunError
from .utils import decode_output

GCC_REPO_URL = 'https://gitee.com/openeuler/gcc.git'
# 克隆成功后才写入的标记:避免上次中断留下的残缺源码树被下一次运行误判为可用
GCC_CLONE_MARK = '.osmts_gcc_clone_complete'

# DejaGnu 的每个 --tool 对应一个测试套件,规模差异很大(按gcc源码树统计):
#   gcc    : 约 6.3 万个测试文件(占约66%)
#   g++    : 约 2.4 万个测试文件(占约25%)
#   gfortran: 约 0.9 万个测试文件(占约9%)
# 三套件合计耗时很长,故支持按工具单独运行,失败时不影响其它工具的结果呈现。
TOOL_RPMS = {'gcc': 'gcc', 'g++': 'gcc-g++', 'gfortran': 'gcc-gfortran'}
ALL_TOOLS = ('gcc', 'g++', 'gfortran')


class DejaGnu:
    """运行 DejaGnu 测试。

    默认跑全部三套件;子类把 tools 收窄为单个工具,便于把耗时的套件拆开单独跑。

    三个工具类共用同一棵 gcc 源码树与同一个测试套件目录
    (/root/osmts_tmp/dejagnu/gcc/gcc/testsuite)。osmts 是顺序执行各测试项的,
    故不存在并发使用的问题;顺带一提,先后运行不同工具时会复用已克隆的源码树。
    """
    tools = ALL_TOOLS

    def __init__(self, **kwargs):
        self.rpms = {'dejagnu'} | {TOOL_RPMS[tool] for tool in self.tools}
        self.path = Path('/root/osmts_tmp/dejagnu')
        self.directory: Path = kwargs.get('saved_directory') / 'dejagnu'
        self.testsuite = self.path / 'gcc/gcc/testsuite/'

    def _ensure_gcc_repo(self):
        """准备gcc源码树。

        三个工具类共用同一棵树,若每次都无条件重克隆,分开跑时会把克隆成本乘三;
        故克隆成功即打标记,后续直接复用(标记缺失说明上次克隆中断,残缺的树不能复用)。
        """
        if (self.path / GCC_CLONE_MARK).exists() and self.testsuite.is_dir():
            print(f"[dejagnu] 复用已就绪的gcc源码:{self.path}")
            return

        shutil.rmtree(self.path, ignore_errors=True)
        # 先克隆到临时目录,成功后再原子改名就位:中断时不会留下"看起来已就绪"的目录
        tmp_path = self.path.parent / (self.path.name + '.tmp')
        shutil.rmtree(tmp_path, ignore_errors=True)
        tmp_path.mkdir(parents=True)
        try:
            subprocess.run(
                f'GIT_TERMINAL_PROMPT=0 git clone {GCC_REPO_URL}',
                cwd=tmp_path,
                shell=True, check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.PIPE,
            )
        except subprocess.CalledProcessError as e:
            shutil.rmtree(tmp_path, ignore_errors=True)
            raise GitCloneError(
                e.returncode, GCC_REPO_URL,
                decode_output(e.stderr) if e.stderr else '',
            )
        (tmp_path / GCC_CLONE_MARK).touch()
        # 就位前再确认一次:Path.rename 在目标已存在时会失败,
        # 而上面的 rmtree 到本行之间若目标被重新创建,就会抛
        # OSError(Directory not empty),白丢一次克隆
        if self.path.exists():
            shutil.rmtree(self.path, ignore_errors=True)
        tmp_path.rename(self.path)

    def pre_test(self):
        if not self.directory.exists():
            self.directory.mkdir(parents=True, exist_ok=True)
        self._ensure_gcc_repo()

    # 解析 dejagnu .sum 文件，将测试结果按 PASS/FAIL/ERROR 等状态分类归纳
    def _parse_sum(self, sum_path):
        # 键必须覆盖正则能匹配到的全部状态:原先正则认识 KFAIL,
        # 字典却没有该键,遇到 KFAIL 行会抛 KeyError 中断整个测试
        result = {'PASS': [], 'FAIL': [], 'UNSUPPORTED': [], 'UNRESOLVED': [],
                  'XFAIL': [], 'XPASS': [], 'WARNING': [], 'ERROR': [],
                  'KFAIL': [], 'other': []}
        summary = {}
        with open(sum_path, errors='ignore') as f:
            for line in f:
                line = line.rstrip()
                m = re.match(r'^(PASS|FAIL|UNSUPPORTED|UNRESOLVED|XFAIL|XPASS|WARNING|ERROR|KFAIL):\s+(.*)$', line)
                if m:
                    status = m.group(1)
                    result[status].append(m.group(2))
                    continue
                sm = re.match(r'^# of (\S+(?:\s+\S+)*)\s+(\d+)$', line)
                if sm:
                    summary[sm.group(1)] = int(sm.group(2))
        return result, summary

    # 将 dejagnu 测试结果按状态优先级写入 Excel 工作表，并附上统计汇总
    def _write_excel(self, results, summary, tool):
        wb = Workbook()
        ws = wb.active
        ws.title = tool

        ws.append(['状态', '测试用例', '详情'])
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 60
        ws.column_dimensions['C'].width = 80

        row = 2
        # 按优先级排列;末尾补上任何未列入的状态,避免将来新增状态时静默丢弃或 KeyError
        status_order = ('FAIL', 'KFAIL', 'XPASS', 'ERROR', 'UNSUPPORTED', 'UNRESOLVED',
                        'WARNING', 'XFAIL', 'other', 'PASS')
        remaining = [s for s in results if s not in status_order]
        for status in status_order + tuple(remaining):
            for case in results.get(status, []):
                case = case.strip()
                m = re.match(r'^(\S+)\s+(.*)$', case)
                test_case = m.group(1) if m else case
                detail = m.group(2).strip() if m else ''
                ws.cell(row=row, column=1, value=status)
                ws.cell(row=row, column=2, value=test_case)
                ws.cell(row=row, column=3, value=detail)
                row += 1

        if summary:
            row += 1
            ws.cell(row=row, column=1, value='汇总统计')
            for key, val in summary.items():
                row += 1
                ws.cell(row=row, column=1, value=key)
                ws.cell(row=row, column=2, value=val)

        wb.save(self.directory / f'{tool}.xlsx')
        print(f"[dejagnu] 已生成 {tool} Excel 报表: {self.directory / f'{tool}.xlsx'}")

    def _run_one_tool(self, tool):
        """运行单个工具的测试套件,返回是否产出了结果文件。"""
        print(f"[dejagnu] 开始运行 {tool} 测试")

        # 先删掉上次的 .sum/.log:否则本次跑不起来时,
        # 残留的旧结果会被误当成"本次成功产出"
        log_file = self.testsuite / f'{tool}.log'
        sum_file = self.testsuite / f'{tool}.sum'
        for stale in (log_file, sum_file):
            try:
                stale.unlink(missing_ok=True)
            except OSError as e:
                print(f"[dejagnu] 无法删除旧结果文件 {stale}: {e}")

        result = subprocess.run(
            f'runtest --tool {tool}',
            cwd=self.testsuite,
            shell=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
        )

        # dejagnu 在找不到全局配置文件时会打印该警告,属正常情况,不算失败
        HARMLESS = "Couldn't find the global config file."
        stderr_lines = [line for line in decode_output(result.stderr).splitlines()
                        if HARMLESS not in line]
        stderr_output = '\n'.join(stderr_lines)

        # 注意:runtest 的返回码语义是"测试结果",不是"命令是否运行成功" ——
        # DejaGnu 的 framework.exp 在有 FAIL/XPASS/KPASS/UNRESOLVED 用例时
        # 会 exit 1(log_and_exit 以 exit_status 退出)。
        # 因此不能把"非0返回码"当成命令失败:否则任何有失败用例的套件
        # (如 gfortran)都会让整个测试报错,连 .sum 里已有的结果都拿不到。
        # 是否真正跑起来以 .sum/.log 为准;返回码只用来判断"跑都没跑起来"。
        if result.returncode != 0:
            print(f"[dejagnu] runtest --tool {tool} 返回码为{result.returncode}"
                  f"(非0通常表示存在FAIL/XPASS等用例,属正常结果)")
        if stderr_output.strip():
            print(f"[{tool}] stderr 警告信息:\n{stderr_output}")

        # 返回码 >2 只可能来自"命令本身没跑起来"(如runtest缺失=127),
        # 此时必然也没有结果文件,给出明确报错而不是静默返回
        if result.returncode > 2:
            raise RunError(
                result.returncode,
                f"dejagnu测试出错.runtest --tool {tool} 未能运行.返回码:"
                f"{result.returncode},报错信息:\n{stderr_output.strip() or '(无)'}",
            )
        if not (log_file.exists() and sum_file.exists()):
            detail = stderr_output.strip() or f"(stderr仅包含无害警告{HARMLESS!r}或为空)"
            raise RunError(
                result.returncode,
                f"dejagnu测试出错.runtest --tool {tool} 未生成{tool}.sum/.log,"
                f"测试未真正运行.返回码:{result.returncode},报错信息:\n{detail}",
            )

        shutil.copy(log_file, self.directory)
        shutil.copy(sum_file, self.directory)

        test_results, summary = self._parse_sum(sum_file)
        self._write_excel(test_results, summary, tool)

        total_pass = len(test_results.get('PASS', []))
        total_fail = len(test_results.get('FAIL', []))
        total_unsupported = len(test_results.get('UNSUPPORTED', []))
        print(f"[dejagnu] {tool} 测试完成: PASS={total_pass} FAIL={total_fail} "
              f"UNSUPPORTED={total_unsupported}")
        return True

    def run_test(self):
        for tool in self.tools:
            self._run_one_tool(tool)

    def run(self):
        print(f'开始进行dejagnu测试,测试项:{list(self.tools)}')
        self.pre_test()
        self.run_test()
        print('dejagnu测试结束')


class DejaGnuGcc(DejaGnu):
    """只运行 gcc 测试套件"""
    tools = ('gcc',)


class DejaGnuGxx(DejaGnu):
    """只运行 g++ 测试套件"""
    tools = ('g++',)


class DejaGnuGfortran(DejaGnu):
    """只运行 gfortran 测试套件"""
    tools = ('gfortran',)